import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

COR = {
    'cabecalho_bg': '1B2A3B', 'cabecalho_fg': 'FFFFFF',
    'ok_bg':        'C6EFCE', 'ok_fg':        '276221',
    'alerta_bg':    'FFEB9C', 'alerta_fg':    '9C6500',
    'reprovado_bg': 'FFC7CE', 'reprovado_fg': '9C0006',
    'perigo_bg':    'FF0000', 'perigo_fg':    'FFFFFF',
    'linha_par':    'F2F2F2', 'linha_impar':  'FFFFFF',
    'resumo_bg':    '2C3E50', 'resumo_fg':    'FFFFFF',
}

def _fill(cor): 
    return PatternFill("solid", fgColor=cor)

def _font(cor, bold=False): 
    return Font(color=cor, bold=bold)

def _border():
    thin = Side(style='thin', color='CCCCCC')
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def _align(h='center', wrap=False):
    return Alignment(horizontal=h, vertical='center', wrap_text=wrap)

def _cor_status(status: str):
    """Retorna (bg, fg) conforme status."""
    m = {
        'OK':       (COR['ok_bg'],       COR['ok_fg']),
        'APROVADO': (COR['ok_bg'],       COR['ok_fg']),
        'ALERTA':   (COR['alerta_bg'],   COR['alerta_fg']),
        'REPROVADO':(COR['reprovado_bg'],COR['reprovado_fg']),
        'PERIGO':   (COR['perigo_bg'],   COR['perigo_fg']),
        'REVISAR':  (COR['reprovado_bg'],COR['reprovado_fg']),
    }
    return m.get(status.upper(), (COR['linha_impar'], '000000'))

def _cor_queda(queda_pct: float, limite: float):
    ratio = queda_pct / limite if limite else 1.0
    if ratio <= 0.75:  return (COR['ok_bg'],    COR['ok_fg'])
    if ratio <= 1.00:  return (COR['alerta_bg'], COR['alerta_fg'])
    return (COR['reprovado_bg'], COR['reprovado_fg'])

import pandas as pd

def exportar_excel(linhas_relatorio, linhas_memorial, caminho_xls, v_linha, tem_trafo):
    df_laudo = pd.DataFrame(linhas_relatorio)
    df_memorial = pd.DataFrame(linhas_memorial)
    
    # Utiliza o motor 'openpyxl' para manter compatibilidade com os seus estilos
    with pd.ExcelWriter(caminho_xls, engine='openpyxl') as writer:
        # Despeja os dados nas duas abas
        df_laudo.to_excel(writer, sheet_name='Laudo Executivo', index=False)
        df_memorial.to_excel(writer, sheet_name='Memorial de Cálculo', index=False)
        
        # Captura os objetos do openpyxl
        wb = writer.book
        ws_laudo = writer.sheets['Laudo Executivo']
        ws_memo = writer.sheets['Memorial de Cálculo']
        
        # Aplica a sua formatação profissional na aba do Laudo
        _formatar_memorial(ws_laudo, df_laudo, v_linha, tem_trafo)
        
        # Aplica a formatação estrutural (bordas, zebra e cabeçalho) na aba do Memorial
        _formatar_memorial(ws_memo, df_memorial, v_linha, tem_trafo)
        
        # Gera a sua aba de Resumo Executivo (ela será inserida como a primeira aba)
        _criar_resumo(wb, df_laudo, v_linha, tem_trafo)


def _formatar_memorial(ws, df: pd.DataFrame, v_ll: float, tem_trafo: bool):
    limite_global = 7.0 if tem_trafo else 5.0

    for cell in ws[1]:
        cell.font      = Font(bold=True, color=COR['cabecalho_fg'])
        cell.fill      = _fill(COR['cabecalho_bg'])
        cell.alignment = _align('center')
        cell.border    = _border()

    col_idx = {cell.value: cell.column for cell in ws[1]}

    for row in ws.iter_rows(min_row=2):
        r = row[0].row
        bg = COR['linha_par'] if r % 2 == 0 else COR['linha_impar']

        for cell in row:
            cell.border    = _border()
            cell.alignment = _align('center')
            cell.fill      = _fill(bg)

        for nome_col in ['Veredito', 'Status Mag.', 'Status Térm.', 'Status Icu']:
            if nome_col in col_idx:
                c = ws.cell(row=r, column=col_idx[nome_col])
                if c.value:
                    bg2, fg2 = _cor_status(str(c.value))
                    c.fill = _fill(bg2)
                    c.font = Font(bold=True, color=fg2)

        for nome_col in ['Queda Acum. (%)']:
            if nome_col in col_idx:
                c = ws.cell(row=r, column=col_idx[nome_col])
                if c.value is not None:
                    bg2, fg2 = _cor_queda(float(c.value), limite_global)
                    c.fill = _fill(bg2)
                    c.font = Font(color=fg2)

        for nome_col in ['Sugestões Técnicas']:
            if nome_col in col_idx:
                c = ws.cell(row=r, column=col_idx[nome_col])
                c.alignment = _align('left', wrap=True)

    LARGURAS = {
        'ID Circuito': 16, 'Quadro Origem': 16, 'Tipo': 14,
        'Fases': 10, 'Comp. (m)': 10, 'Cabo (mm²)': 11, 'Neutro (mm²)': 12,
        'Ib (A)': 9, 'FP': 7, 'Disjuntor': 14, 'Isolação': 10,
        'Tensão (V)': 12, 'Queda Acum. (%)': 15, 'Queda Trecho (%)': 16,
        'Icc Mín. (A)': 13, 'Icc Máx. (A)': 13, 'Ip Pico (A)': 12,
        'Status Mag.': 13, 'Status Térm.': 13, 'Status Icu': 11,
        'Veredito': 12, 'Sugestões Técnicas': 55,
    }
    for col in ws.columns:
        letra = col[0].column_letter
        nome  = ws.cell(row=1, column=col[0].column).value
        ws.column_dimensions[letra].width = LARGURAS.get(nome, 14)

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

def _criar_resumo(wb, df: pd.DataFrame, v_ll: float, tem_trafo: bool):
    ws = wb.create_sheet('Resumo Executivo', 0)  # primeira aba
    limite_global = 7.0 if tem_trafo else 5.0

    total        = len(df)
    aprovados    = (df['Veredito'] == 'APROVADO').sum()
    reprovados   = total - aprovados
    pior_queda   = df['Queda Acum. (%)'].max()
    icc_min_min  = df['Icc Mín. (A)'].min()
    icc_max_max  = df['Icc Máx. (A)'].max()
    taxa_aprov   = (aprovados / total * 100) if total else 0

    h_fill  = _fill(COR['resumo_bg'])
    h_font  = Font(bold=True, color=COR['resumo_fg'], size=11)
    v_font  = Font(size=11)
    b       = _border()

    def escrever(linha, col, valor, negrito=False, bg=None, fg='000000', align='left'):
        c = ws.cell(row=linha, column=col, value=valor)
        c.font      = Font(bold=negrito, color=fg, size=11)
        c.border    = b
        c.alignment = _align(align)
        if bg:
            c.fill = _fill(bg)

    ws.merge_cells('A1:D1')
    t = ws['A1']
    t.value     = 'RESUMO EXECUTIVO — MOTOR NODAL V4.5'
    t.font      = Font(bold=True, color='FFFFFF', size=14)
    t.fill      = _fill(COR['cabecalho_bg'])
    t.alignment = _align('center')

    ws.merge_cells('A2:D2')
    ws['A2'].value     = f'Sistema {v_ll:.0f} V  |  Limite global de queda: {limite_global:.0f}%  |  {"Trafo próprio" if tem_trafo else "Alimentado pela concessionária"}'
    ws['A2'].font      = Font(italic=True, color='555555', size=10)
    ws['A2'].alignment = _align('center')

    dados = [
        ('Total de circuitos',       total,               None,                  '000000'),
        ('Aprovados',                aprovados,           COR['ok_bg'],          COR['ok_fg']),
        ('Em revisão / reprovados',  reprovados,          COR['reprovado_bg'],   COR['reprovado_fg']),
        ('Taxa de aprovação (%)',     f'{taxa_aprov:.1f}%',
                                     COR['ok_bg'] if taxa_aprov == 100 else COR['alerta_bg'],
                                     COR['ok_fg'] if taxa_aprov == 100 else COR['alerta_fg']),
        ('Pior queda acumulada (%)', f'{pior_queda:.2f}%',
                                     *_cor_queda(pior_queda, limite_global)),
        ('Icc mínimo do sistema (A)',f'{icc_min_min:.1f}', None, '000000'),
        ('Icc máximo do sistema (A)',f'{icc_max_max:.1f}', None, '000000'),
    ]

    for i, (label, valor, bg, fg) in enumerate(dados, start=4):
        escrever(i, 1, label,  negrito=True,  bg='F5F5F5')
        escrever(i, 2, valor,  negrito=False, bg=bg, fg=fg, align='center')
        ws.row_dimensions[i].height = 20

    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 10
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 18